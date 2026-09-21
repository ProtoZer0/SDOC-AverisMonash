"""Format dispatch: bytes in, a parsed document out.

Every failure mode is an explicit outcome, never an exception that escapes:
the pipeline has to be able to say WHY a document could not be read.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import PurePosixPath

from parsers import docx as docx_reader
from parsers import pdf as pdf_reader
from parsers import txt as txt_reader
from parsers import xlsx as xlsx_reader
from parsers.common import HitBag, sha256
from parsers.doctype import detect
from pipeline.schemas import DocFormat, DocKind, ShipmentFields


@dataclass
class ParsedDocument:
    fmt: DocFormat
    readable: bool
    fields: ShipmentFields
    bag: HitBag | None
    full_text: str
    text_sha256: str | None
    page_count: int | None
    detected_kind: DocKind
    needs_ocr: bool = False
    parse_error: str | None = None
    ocr_confidence: float | None = None


_EXT_TO_FMT: dict[str, DocFormat] = {
    ".txt": "txt",
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
}


def format_of(path: str) -> DocFormat:
    return _EXT_TO_FMT.get(PurePosixPath(path).suffix.lower(), "txt")


def _empty(fmt: DocFormat, error: str, *, needs_ocr: bool = False, pages: int | None = None) -> ParsedDocument:
    return ParsedDocument(
        fmt=fmt,
        readable=False,
        fields=ShipmentFields(),
        bag=None,
        full_text="",
        text_sha256=None,
        page_count=pages,
        detected_kind="UNKNOWN",
        needs_ocr=needs_ocr,
        parse_error=error,
    )


def read(path: str, data: bytes) -> ParsedDocument:
    fmt = format_of(path)
    sheet_names: list[str] = []

    try:
        if fmt == "pdf":
            try:
                bag, text, pages = pdf_reader.read(data)
            except pdf_reader.ImageOnlyPDF as exc:
                return _empty("scan_pdf", "image-only pdf", needs_ocr=True, pages=exc.page_count)
            except pdf_reader.UnreadablePDF as exc:
                return _empty("pdf", f"unreadable pdf: {exc}")
        elif fmt == "docx":
            bag, text, pages = docx_reader.read(data)
        elif fmt == "xlsx":
            bag, text, pages = xlsx_reader.read(data)
            sheet_names = [ln.split(" | ")[0] for ln in []]  # populated below
        else:
            bag, text, pages = txt_reader.read(data)
    except Exception as exc:  # noqa: BLE001 — a bad file is data, not a crash
        return _empty(fmt, f"{type(exc).__name__}: {exc}")

    if fmt == "xlsx":
        sheet_names = _sheet_names(data)

    fields = bag.to_fields()
    kind = detect(text, sheet_names=sheet_names, fields_found=fields.found_count())

    return ParsedDocument(
        fmt=fmt,
        readable=True,
        fields=fields,
        bag=bag,
        full_text=text,
        text_sha256=sha256(text),
        page_count=pages,
        detected_kind=kind,
    )


def _sheet_names(data: bytes) -> list[str]:
    import io

    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    except Exception:  # noqa: BLE001
        return []
